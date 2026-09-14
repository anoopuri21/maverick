"""Editable DOCX renderer for the Bachelors page plan (python-docx)."""
import gen_page_plan_keywords as C

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "landing-page/bachelors/05-page-plan-keywords.docx"

NAVY = RGBColor(0x07, 0x14, 0x44)
RED = RGBColor(0xB2, 0x02, 0x02)
INK = RGBColor(0x1C, 0x1E, 0x26)
GREY = RGBColor(0x5A, 0x60, 0x70)


def set_base_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, bold in (("Heading 1", 15, NAVY, True),
                                    ("Heading 2", 12.5, NAVY, True),
                                    ("Heading 3", 11.5, NAVY, True)):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = bold
        h.font.color.rgb = color
        h.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        h.paragraph_format.space_after = Pt(5)


def para(doc, text, size=11, bold=False, italic=False, color=None, space_after=6, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def bullet(doc, text, lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if lead:
        r = p.add_run(lead + " ")
        r.font.bold = True
        r.font.color.rgb = NAVY
    p.add_run(text)
    p.paragraph_format.space_after = Pt(4)
    return p


def chart_table(doc, chart):
    """Chart data as an editable table (charts on the live page are drawn in brand colours)."""
    para(doc, chart["caption"], size=10.5, bold=True, color=NAVY, space_after=3)
    if chart["kind"] == "pie":
        rows = [(label, f"{pct}%") for label, pct, _ in chart["slices"]]
    else:
        rows = []
        for label, lo, hi in chart["bars"]:
            val = "added once your fees are confirmed" if lo is None else (
                f"AED {lo:g}k\u2013{hi:g}k" if hi != lo else f"AED {lo:g}k")
            rows.append((label, val))
    t = doc.add_table(rows=len(rows) + 1, cols=2)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text = "Item"
    hdr[1].text = "Value"
    for cell in hdr:
        for pr in cell.paragraphs:
            for r in pr.runs:
                r.font.bold = True
                r.font.color.rgb = NAVY
    for i, (a, bv) in enumerate(rows, 1):
        t.rows[i].cells[0].text = a
        t.rows[i].cells[1].text = bv
    para(doc, chart["source"], size=9, color=GREY, space_after=2)
    para(doc, "On the live page this appears as a chart in the brand colours, with the source line visible.",
         size=9, italic=True, color=GREY, space_after=10)


def block15(doc, name, cat, heading, desc, scope, chart):
    doc.add_heading(name, level=3)
    para(doc, f"Category: {cat}", size=9, bold=True, color=GREY, space_after=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r1 = p.add_run("Heading: ")
    r1.font.bold = True
    r1.font.color.rgb = NAVY
    p.add_run(f"\u201c{heading}\u201d")
    para(doc, desc, size=10, italic=True, space_after=3)
    para(doc, scope, size=10.5, space_after=8)
    if chart:
        chart_table(doc, chart)


def build_docx():
    C.assert_clean()
    doc = Document()
    for section in doc.sections:  # clean sheet, no header/footer content
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.1)
        section.right_margin = Cm(2.1)
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True
    set_base_styles(doc)

    # title block
    para(doc, C.TITLE_1, size=20, bold=True, color=NAVY, space_after=2)
    para(doc, C.TITLE_2, size=14, bold=True, color=RED, space_after=8)
    para(doc, C.ORG_LINE, size=9, color=GREY, space_after=10)
    para(doc, C.SUBTITLE, size=11, space_after=12)

    # contents
    doc.add_heading("What is inside", level=1)
    for i, line in enumerate(C.CONTENTS, 1):
        para(doc, f"{i}. {line}", space_after=2)

    # 1
    doc.add_heading("1. The page in one minute", level=1)
    para(doc, C.S1_INTRO)
    para(doc, "The page offers four study routes, all from our own catalogue.")
    for title, body in C.ROUTES:
        para(doc, title, bold=True, color=NAVY, space_after=2)
        para(doc, body, space_after=8)
    para(doc, C.S1_OUTRO)

    # 2
    doc.add_heading("2. Who the page is for", level=1)
    para(doc, C.S2_INTRO)
    for title, body in C.PERSONAS:
        para(doc, title, bold=True, color=NAVY, space_after=2)
        para(doc, body, space_after=8)

    # 3
    doc.add_heading("3. The market at a glance", level=1)
    para(doc, C.S3_INTRO)
    for _, body in C.S3_BULLETS:
        bullet(doc, body)

    # 4
    doc.add_heading("4. What the page promises", level=1)
    para(doc, C.S4_INTRO)
    for lead, body in C.PROMISES:
        bullet(doc, body, lead=lead)

    # 5
    doc.add_heading("5. The page, section by section", level=1)
    para(doc, C.S5_INTRO)
    for args in C.BLOCKS:
        block15(doc, *args)

    # 6
    doc.add_heading("6. How the page will be found online", level=1)
    para(doc, C.S6_INTRO)
    for title, body in C.S6_PARTS:
        doc.add_heading(title, level=2)
        para(doc, body)

    # 7
    doc.add_heading("7. The search phrases (keywords)", level=1)
    para(doc, C.S7_INTRO)
    doc.add_heading("Tier 1: the page's core search phrases", level=2)
    for code, phrase, body in C.TIER1:
        para(doc, f"{code}  {phrase}", bold=True, color=NAVY, space_after=2)
        para(doc, body, size=10.5, space_after=8)
    doc.add_heading("Tier 2: broader reach phrases (woven in, not headlined)", level=2)
    for code, phrase, body in C.TIER2:
        para(doc, f"{code}  {phrase}", bold=True, color=NAVY, space_after=2)
        para(doc, body, size=10.5, space_after=8)
    para(doc, C.SUPPORTING)

    doc.save(OUT)
    print("wrote", OUT)
