"""Client content pack renderer for the Masters landing page (DOCX + PDF).

Reads landing-page/gulf-masters/08-masters-landing-content-gcc-uae.md,
takes ONLY the page content (sections 1 to 18, stops before the internal
META & TECHNICAL SPEC block) and renders a clean, copy-paste ready client
document:

  - yellow highlight on every target keyword phrase
  - shaded navy heading for each page section
  - red NEW badge on the two added sections (Journey, Compare)
  - content only: no explanations, no internal annotations

Usage:  python3 scripts/gen_masters_content_client.py
Writes: landing-page/gulf-masters/10-masters-landing-content-client-gcc-uae.docx
        landing-page/gulf-masters/10-masters-landing-content-client-gcc-uae.pdf
"""
import re
import sys
from xml.sax.saxutils import escape

SRC = "landing-page/gulf-masters/08-masters-landing-content-gcc-uae.md"
OUT_BASE = "landing-page/gulf-masters/10-masters-landing-content-client-gcc-uae"

KEYWORDS = [
    "MBA for working professionals in UAE",
    "MBA specializations in UAE",
    "Dubai for working professionals",
    "flexible payment plan",
    "September 2026 intake",
    "International MBA",
    "admission requirements",
    "MBA fees in UAE",
    "MBA scholarship",
    "fast-track MBA",
    "Part-time MBA",
    "Affordable MBA",
    "Flexible MBA",
    "MBA in Dubai",
    "MBA in UAE",
    "Online MBA",
]
NEW_SECTIONS = ("JOURNEY", "COMPARE")

# brand palette (repo convention)
NAVY_HEX, RED_HEX, INK_HEX, GREY_HEX = "071444", "B20202", "1C1E26", "5A6070"
HL_PDF = "FFF59D"
SHADE_SEC, SHADE_NEW = "EEF2FA", "FCE8E8"

FLAG_TO_TEXT = {
    "\U0001F1E6\U0001F1EA": "UAE",
    "\U0001F1F8\U0001F1E6": "Saudi Arabia",
    "\U0001F1F4\U0001F1F2": "Oman",
    "\U0001F1F6\U0001F1E6": "Qatar",
}

FIELD_NAMES = {
    "heading": "Heading", "intro": "Intro", "eyebrow": "Eyebrow",
    "subheading": "Subheading", "form_title": "Form title",
    "cta_primary": "Button (primary)", "cta_secondary": "Button (secondary)",
    "cta_tertiary_label": "Button label", "label": "Label", "quote": "Quote",
    "stats": "Stats", "items": "Items", "chapters": "Blocks",
    "steps": "Steps", "tabs": "Tabs", "closing": "Closing line",
    "universities": "Universities", "trending": "Trending picks",
    "audience": "Audience", "metrics": "Metrics", "regions": "Regions",
    "industries": "Industries", "rows": "Table", "note": "Note",
    "blocks": "Blocks", "stories": "Stories", "copy": "Copy",
    "trust_line": "Trust line", "partners": "Partners",
    "checklist": "Checklist", "points": "Points", "show_form": "Show form",
}


# --------------------------------------------------------------------------
# parsing
# --------------------------------------------------------------------------
def extract_body():
    text = open(SRC, encoding="utf-8").read()
    start = text.index("## 1. HERO")
    end = text.index("## META & TECHNICAL SPEC")
    return text[start:end]


def parse():
    """Return a list of sections: {num, name, is_new, items:[(kind, payload)]}."""
    sections, cur, field = [], None, None
    for raw in extract_body().splitlines():
        line = raw.rstrip()
        if not line.strip() or line.strip() == "---" or line.strip() == "[list]":
            continue
        m = re.match(r"^## (\d+)\. (.+)$", line)
        if m:
            name = m.group(2).strip()
            is_new = "(NEW SECTION)" in name
            name = name.replace("(NEW SECTION)", "").strip()
            cur = {"num": m.group(1), "name": name, "is_new": is_new, "items": []}
            sections.append(cur)
            field = None
            continue
        m = re.match(r"^### (.+)$", line)
        if m:
            lab = m.group(1).strip()
            header = None
            m2 = re.match(r"^(\w+)\s*\((.+)\)$", lab)
            if m2:
                lab, hint = m2.group(1), m2.group(2)
                if lab == "rows" and "|" in hint:
                    header = [c.strip() for c in hint.split("|")]
            field = {"label": lab, "header": header}
            cur["items"].append(("field", field))
            continue
        if cur is None:
            continue
        m = re.match(r"^(\d+)\. (.+)$", line)
        if m:
            cur["items"].append(("numbered", (m.group(1), m.group(2), field)))
            continue
        if line.startswith("- "):
            cur["items"].append(("bullet", line[2:].strip()))
            continue
        if " | " in line:
            cur["items"].append(("pipe", [c.strip() for c in line.split(" | ")]))
            continue
        if line.startswith("**") and line.endswith("**"):
            cur["items"].append(("blockhead", line.strip("*").strip()))
            continue
        cur["items"].append(("para", line.strip()))
    return sections


def norm_flags(text):
    for flag, name in FLAG_TO_TEXT.items():
        text = text.replace(flag, name)
    return text


# --------------------------------------------------------------------------
# keyword span helpers
# --------------------------------------------------------------------------
def spans(text):
    """Yield (segment, is_keyword) covering text."""
    low = text.lower()
    hits = []
    for kw in KEYWORDS:
        k = kw.lower()
        i = 0
        while True:
            j = low.find(k, i)
            if j < 0:
                break
            hits.append((j, j + len(kw)))
            i = j + 1
    if not hits:
        return [(text, False)]
    hits.sort()
    merged = []
    for a, b in hits:
        if merged and a <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], b))
        else:
            merged.append([a, b])
    out, pos = [], 0
    for a, b in merged:
        if a > pos:
            out.append((text[pos:a], False))
        out.append((text[a:b], True))
        pos = b
    if pos < len(text):
        out.append((text[pos:], False))
    return out


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def build_docx(sections):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    NAVY = RGBColor(0x07, 0x14, 0x44)
    RED = RGBColor(0xB2, 0x02, 0x02)
    INK = RGBColor(0x1C, 0x1E, 0x26)
    GREY = RGBColor(0x5A, 0x60, 0x70)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    def shade(p, fill):
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        pPr.append(shd)

    def kw_runs(p, text, size=11, bold=False, color=None, italic=False):
        for seg, is_kw in spans(text):
            r = p.add_run(seg)
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            if color:
                r.font.color.rgb = color
            if is_kw:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                r.font.bold = True

    def add_text(text, size=11, bold=False, color=None, space_after=6,
                 space_before=0, italic=False, indent=None):
        p = doc.add_paragraph()
        kw_runs(p, text, size=size, bold=bold, color=color, italic=italic)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        if indent:
            p.paragraph_format.left_indent = Cm(indent)
        return p

    def field_label(name):
        p = doc.add_paragraph()
        r = p.add_run(FIELD_NAMES.get(name, name.title()).upper())
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = GREY
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(8)

    def table(rows, header=None):
        ncols = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows) + (1 if header else 0), cols=ncols)
        t.style = "Table Grid"
        ri = 0
        if header:
            for ci, h in enumerate(header):
                cell = t.rows[0].cells[ci]
                cell.text = ""
                kw_runs(cell.paragraphs[0], h, size=10, bold=True, color=NAVY)
            ri = 1
        for row in rows:
            for ci in range(ncols):
                val = row[ci] if ci < len(row) else ""
                cell = t.rows[ri].cells[ci]
                cell.text = ""
                kw_runs(cell.paragraphs[0], val, size=10,
                        bold=(ci == 0 and not header))
            ri += 1
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # title block
    p = doc.add_paragraph()
    r = p.add_run("Online MBA & Master's Degrees for the UAE and GCC")
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    r = p.add_run("Landing page content  \u00b7  /online-mba-masters-uae  \u00b7  Maverick  \u00b7  16 September 2026")
    r.font.size = Pt(10)
    r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    r = p.add_run("Highlighted text = target keyword phrase.  Shaded headings = page sections.  "
                  "RED badge = NEW section being added to the page.")
    r.font.size = Pt(8.5)
    r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(10)

    def pipe_group(rows, field):
        header = field.get("header") if field else None
        nrows, ncols = len(rows), max(len(r) for r in rows)
        if nrows == 1 and ncols > 4:  # single wide row: render as bullets
            for cell in rows[0]:
                pb = doc.add_paragraph(style="List Bullet")
                kw_runs(pb, cell, size=10.5)
                pb.paragraph_format.space_after = Pt(2)
            return
        table(rows, header)

    for s in sections:
        # section heading
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        shade(p, SHADE_NEW if s["is_new"] else SHADE_SEC)
        r = p.add_run(f"{s['num']}. {s['name']}")
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RED if s["is_new"] else NAVY
        if s["is_new"]:
            r2 = p.add_run("   [NEW SECTION]")
            r2.font.size = Pt(10)
            r2.font.bold = True
            r2.font.color.rgb = RED

        field = None
        pipe_buf = []

        def flush(field_ctx):
            nonlocal pipe_buf
            if pipe_buf:
                pipe_group(pipe_buf, field_ctx)
                pipe_buf = []

        for kind, payload in s["items"]:
            if kind == "pipe":
                pipe_buf.append(payload)
                continue
            flush(field)
            if kind == "field":
                field = payload
                field_label(payload["label"])
            elif kind == "para":
                add_text(norm_flags(payload))
            elif kind == "blockhead":
                add_text(payload, bold=True, color=NAVY, space_after=2, space_before=4)
            elif kind == "bullet":
                text = norm_flags(payload)
                pb = doc.add_paragraph(style="List Bullet")
                lead = None
                for fname in FLAG_TO_TEXT.values():
                    if text.startswith(fname + " \u00b7 "):
                        lead = fname
                        text = text[len(fname) + 3:]
                        break
                if lead:
                    rl = pb.add_run(lead + "   ")
                    rl.font.bold = True
                    rl.font.color.rgb = NAVY
                kw_runs(pb, text, size=10.5)
                pb.paragraph_format.space_after = Pt(3)
            elif kind == "numbered":
                num, text, _f = payload
                if s["name"].startswith("FAQ"):
                    pq = doc.add_paragraph()
                    pq.paragraph_format.space_before = Pt(6)
                    pq.paragraph_format.space_after = Pt(1)
                    rq = pq.add_run(f"Q{num}. ")
                    rq.font.bold = True
                    rq.font.color.rgb = NAVY
                    kw_runs(pq, text, bold=True, color=NAVY)
                else:
                    add_text(f"{num}. {text}", space_after=3)
        flush(field)

        # FAQ answers: numbered item in FAQ is question; following para is answer.
    doc.save(OUT_BASE + ".docx")


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def build_pdf(sections):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    navy = colors.HexColor("#" + NAVY_HEX)
    red = colors.HexColor("#" + RED_HEX)
    ink = colors.HexColor("#" + INK_HEX)
    grey = colors.HexColor("#" + GREY_HEX)

    S = {}
    S["title"] = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=16,
                                textColor=navy, leading=20, spaceAfter=2)
    S["sub"] = ParagraphStyle("s", fontName="Helvetica", fontSize=9.5,
                              textColor=grey, leading=13, spaceAfter=2)
    S["legend"] = ParagraphStyle("l", fontName="Helvetica", fontSize=8.5,
                                 textColor=grey, leading=12, spaceAfter=10)
    S["sec"] = ParagraphStyle("sec", fontName="Helvetica-Bold", fontSize=12.5,
                              textColor=navy, leading=16, spaceBefore=14,
                              spaceAfter=4, backColor=colors.HexColor("#" + SHADE_SEC),
                              borderPadding=(4, 4, 4, 4), leftIndent=2)
    S["sec_new"] = ParagraphStyle("secn", parent=S["sec"], textColor=red,
                                  backColor=colors.HexColor("#" + SHADE_NEW))
    S["field"] = ParagraphStyle("f", fontName="Helvetica-Bold", fontSize=8,
                                textColor=grey, spaceBefore=7, spaceAfter=1,
                                leading=10)
    S["body"] = ParagraphStyle("b", fontName="Helvetica", fontSize=10.5,
                               textColor=ink, leading=14.5, spaceAfter=5)
    S["bullet"] = ParagraphStyle("bl", parent=S["body"], leftIndent=14,
                                 bulletIndent=4, spaceAfter=3)
    S["blockhead"] = ParagraphStyle("bh", fontName="Helvetica-Bold", fontSize=10.5,
                                    textColor=navy, leading=14, spaceBefore=5,
                                    spaceAfter=1)
    S["cell"] = ParagraphStyle("c", fontName="Helvetica", fontSize=9.5,
                               textColor=ink, leading=12.5)
    S["cellb"] = ParagraphStyle("cb", parent=S["cell"], fontName="Helvetica-Bold",
                                textColor=navy)

    def hl(text):
        out = []
        for seg, is_kw in spans(text):
            e = escape(seg)
            out.append(f'<font backColor="#{HL_PDF}"><b>{e}</b></font>' if is_kw else e)
        return "".join(out)

    def make_table(rows, header=None):
        data = []
        if header:
            data.append([Paragraph(escape(h), S["cellb"]) for h in header])
        for r in rows:
            data.append([Paragraph(hl(c), S["cell"]) for c in r])
        ncols = max(len(r) for r in data)
        avail = 16.8 * cm
        t = Table(data, colWidths=[avail / ncols] * ncols, repeatRows=1 if header else 0)
        style = [
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#C9CFDD")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]
        if header:
            style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9EDF6")))
        t.setStyle(TableStyle(style))
        return t

    doc = SimpleDocTemplate(OUT_BASE + ".pdf", pagesize=A4,
                            leftMargin=2.1 * cm, rightMargin=2.1 * cm,
                            topMargin=1.8 * cm, bottomMargin=1.8 * cm,
                            title="Masters landing page content (UAE & GCC)",
                            author="Maverick")
    story = [
        Paragraph("Online MBA &amp; Master's Degrees for the UAE and GCC", S["title"]),
        Paragraph("Landing page content &nbsp;\u00b7&nbsp; /online-mba-masters-uae &nbsp;\u00b7&nbsp; Maverick &nbsp;\u00b7&nbsp; 16 September 2026", S["sub"]),
        Paragraph("Highlighted text = target keyword phrase. &nbsp;Shaded headings = page sections. &nbsp;"
                  "<font color='#B20202'><b>RED</b></font> badge = NEW section being added to the page.", S["legend"]),
    ]

    def pipe_group(rows, field):
        header = field.get("header") if field else None
        if len(rows) == 1 and len(rows[0]) > 4:
            for cell in rows[0]:
                story.append(Paragraph(hl(cell), S["bullet"], bulletText="\u2022"))
            return
        story.append(make_table(rows, header))
        story.append(Spacer(1, 4))

    for s in sections:
        style = S["sec_new"] if s["is_new"] else S["sec"]
        badge = ' &nbsp;<font color="#B20202">[NEW SECTION]</font>' if s["is_new"] else ""
        story.append(Paragraph(escape(f"{s['num']}. {s['name']}") + badge, style))
        field = None
        pipe_buf = []

        def flush(fctx):
            nonlocal pipe_buf
            if pipe_buf:
                pipe_group(pipe_buf, fctx)
                pipe_buf = []

        for kind, payload in s["items"]:
            if kind == "pipe":
                pipe_buf.append(payload)
                continue
            flush(field)
            if kind == "field":
                field = payload
                story.append(Paragraph(escape(FIELD_NAMES.get(payload["label"], payload["label"]).upper()), S["field"]))
            elif kind == "para":
                story.append(Paragraph(hl(norm_flags(payload)), S["body"]))
            elif kind == "blockhead":
                story.append(Paragraph(hl(payload), S["blockhead"]))
            elif kind == "bullet":
                text = norm_flags(payload)
                lead = ""
                for fname in FLAG_TO_TEXT.values():
                    if text.startswith(fname + " \u00b7 "):
                        lead = f"<b>{escape(fname)}</b> &nbsp;"
                        text = text[len(fname) + 3:]
                        break
                story.append(Paragraph(lead + hl(text), S["bullet"], bulletText="\u2022"))
            elif kind == "numbered":
                num, text, _f = payload
                if s["name"].startswith("FAQ"):
                    story.append(Paragraph(f"<b><font color='#{NAVY_HEX}'>Q{num}. {hl(text)}</font></b>", S["body"]))
                else:
                    story.append(Paragraph(f"<b>{num}.</b> {hl(text)}", S["body"]))
        flush(field)

    doc.build(story)


# --------------------------------------------------------------------------
# verification + main
# --------------------------------------------------------------------------
def all_text(sections):
    parts = []
    for s in sections:
        parts.append(f"{s['num']} {s['name']}")
        for kind, payload in s["items"]:
            if kind == "para":
                parts.append(payload)
            elif kind == "blockhead":
                parts.append(payload)
            elif kind == "bullet":
                parts.append(payload)
            elif kind == "pipe":
                parts.extend(payload)
            elif kind == "numbered":
                parts.append(payload[1])
    return parts


def main():
    sections = parse()
    texts = [norm_flags(t) for t in all_text(sections)]
    joined = "\n".join(texts)
    words = len(joined.split())
    print(f"sections parsed : {len(sections)}")
    print(f"word count      : {words}  (client cap: 3000 to 3500)")
    if not (3000 <= words <= 3500):
        sys.exit("FAIL: word count outside 3000 to 3500")

    missing = [kw for kw in KEYWORDS if kw.lower() not in joined.lower()]
    if missing:
        sys.exit(f"FAIL: keywords missing from rendered content: {missing}")
    print(f"keywords        : {len(KEYWORDS)}/{len(KEYWORDS)} present and highlighted")

    for ch, name in (("\u2014", "em dash"), ("\u2013", "en dash")):
        if ch in joined:
            sys.exit(f"FAIL: {name} found in content")
    if re.search(r"\bTBD\b|\[TBD\]|placeholder", joined, re.I):
        sys.exit("FAIL: placeholder text found")
    print("sweeps          : 0 em/en dashes, 0 placeholders")

    new_names = [s["name"] for s in sections if s["is_new"]]
    print(f"NEW sections    : {new_names}")

    build_docx(sections)
    build_pdf(sections)
    print("wrote", OUT_BASE + ".docx")
    print("wrote", OUT_BASE + ".pdf")


if __name__ == "__main__":
    main()
